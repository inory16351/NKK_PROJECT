"""초필살기 새 글 번역 연결 (2026-10-18): ① 스트링 테이블 '번역 필요' 중 아래 EN 에 있는 글에 영어 채움 ② 유니티에 붙일 C# 생성
   (Game: HUD/SuperUlt 아래 한글 TMP 에 LocText · Lobby: 쳇바퀴 훈련 초필살기 카드 TMP 에 LocText + TrainPage.superPathTag 를 LocScene fields 에)
순서: 유니티에서 씬 저장 → python Tools/import_scene_strings.py → python Tools/hook_superult_strings.py <game.cs> <lobby.cs>
      → gen_string_table · xlsx2json → 각 .cs 를 그 씬을 연 상태에서 MCP execute_code"""
import json, os, re, sys
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
EN = {
    '필살기를 {n}번 끝까지 쓰면 버튼이 떠요 · 종마다 한 판에 한 번': 'Use its Ultimate {n} times to unlock the button · once per kind per run',
    ' + 초필살기': ' + Super Ultimate',
    'Lv {lv}에 해금': 'Unlocks at Lv {lv}',          # 이미 있는 글 (키 연결용)
    '초필살기 · {name}': 'Super Ultimate · {name}',
    '초필살기!!': 'SUPER!!',
    '초필살기 이름': 'Super Ultimate',
    '{name}의 초필살기': "{name}'s Super Ultimate",
    '{n}명': '{n} subs',
    '※ 이번 달 전기세는 연구소가 냅니다': "※ This month's power bill goes to the Lab",
    '※ 줄 서기 3시간, 퍼레이드 3초': '※ 3 hours in line, 3 seconds of parade',
    '※ 대관식 비용은 연구소 법인카드로': "※ Coronation billed to the Lab's company card",
    '※ 자외선 지수: 매우 위험': '※ UV index: EXTREME',
    '※ 원탁이라 상석 다툼 없음 (아마도)': '※ Round table, so no fighting over seats (probably)',
    '※ 소원 추가 요청은 반려되었습니다': '※ Request for more wishes: DENIED',
    '※ 오늘 금화 비는 술탄이 쏩니다': "※ Today's gold rain is on the Sultan",
    '※ 소는 촬영 후 무사히 돌려보냈습니다': '※ The cow was safely returned after filming',
    '※ 구독 · 좋아요 · 알림 설정까지!': '※ Like, subscribe and hit the bell!',
    '※ 다람쥐 헌 쳇바퀴에 타고파~': '※ Squirrel wants a spin on the old wheel~',
    '※ 앵콜 요청은 마음속으로만': '※ Please keep encore requests in your heart',
    '※ 익힘 정도: 완벽하게 이븐함': '※ Doneness: perfectly even',
    '※ 범인은 이 안에 있다!': '※ The culprit is among us!',
    '※ 쭈욱~ 늘어나는 건 국룰': '※ The cheese pull is mandatory',
    '※ 시전 후 MP 0 · 이동 불가': '※ After casting: MP 0 · cannot move',
    '※ 얼어 죽어도 아이스 아메리카노': '※ Iced Americano, even if I freeze',
    '※ 시간 외 근무 수당은 청구하지 않습니다': '※ No overtime pay will be claimed',     # 메이드 쥐 (2026-10-20)
}
TEMPLATE = re.compile(r'\{\w+\}')
CODE_SET = ('Title', 'Line')          # 코드가 테이블 글로 채우는 TMP (틀만)


def fill():
    p = os.path.join(ROOT, 'Data_Table', '스트링 테이블.xlsx')
    wb = openpyxl.load_workbook(p); ws = wb['String']
    k = [c.value for c in ws[2]]; ko, en, st = k.index('ko') + 1, k.index('en') + 1, k.index('status') + 1
    keymap, n = {}, 0
    for r in range(4, ws.max_row + 1):
        s = ws.cell(r, ko).value
        if s is None: continue
        keymap.setdefault(s, int(ws.cell(r, 1).value))
        if ws.cell(r, st).value in ('번역 필요', '한글 바뀜') and s in EN:
            ws.cell(r, en).value = EN[s]; ws.cell(r, st).value = None; n += 1
    wb.save(p)
    print('영어 채움', n)
    return keymap


def cs_map(keymap):
    items = [(s, keymap[s]) for s in EN if s in keymap]
    return 'var km = new System.Collections.Generic.Dictionary<string, int>();\n' + '\n'.join(f'km["{s.replace(chr(34), chr(92) + chr(34))}"] = {v};' for s, v in items)


def game_cs(keymap):
    return cs_map(keymap) + '''
var root = GameObject.Find("HUD").transform.Find("SuperUlt");
int n = 0;
foreach (var t in root.GetComponentsInChildren<TMPro.TMP_Text>(true))
{
  if (!km.ContainsKey(t.text)) continue;
  var lt = t.GetComponent<NKK.LocText>(); if (!lt) lt = t.gameObject.AddComponent<NKK.LocText>();
  lt.key = km[t.text]; lt.template = t.text.Contains("{") || t.name == "Title" || t.name == "Line"; EditorUtility.SetDirty(lt); n++;
}
UnityEditor.SceneManagement.EditorSceneManager.MarkSceneDirty(root.gameObject.scene);
UnityEditor.SceneManagement.EditorSceneManager.SaveScene(root.gameObject.scene);
return "LocText " + n;'''


def lobby_cs(keymap):
    return cs_map(keymap) + '''
var tp = UnityEngine.Object.FindFirstObjectByType<NKK.Lobby.TrainPage>(FindObjectsInactive.Include);
int n = 0;
foreach (var t in tp.superCard.GetComponentsInChildren<TMPro.TMP_Text>(true))
{
  if (!km.ContainsKey(t.text)) continue;
  var lt = t.GetComponent<NKK.LocText>(); if (!lt) lt = t.gameObject.AddComponent<NKK.LocText>();
  lt.key = km[t.text]; lt.template = t.text.Contains("{"); EditorUtility.SetDirty(lt); n++;
}
var ls = UnityEngine.Object.FindFirstObjectByType<NKK.LocScene>(FindObjectsInactive.Include);
ls.fields.RemoveAll(x => x.target == tp && x.path == "superPathTag");
if (km.ContainsKey(tp.superPathTag)) { var f = new NKK.LocScene.Field(); f.target = tp; f.path = "superPathTag"; f.key = km[tp.superPathTag]; ls.fields.Add(f); n++; }
EditorUtility.SetDirty(ls);
UnityEditor.SceneManagement.EditorSceneManager.MarkSceneDirty(tp.gameObject.scene);
UnityEditor.SceneManagement.EditorSceneManager.SaveScene(tp.gameObject.scene);
return "Lobby LocText " + n;'''


if __name__ == '__main__':
    km = fill()
    open(sys.argv[1], 'w', encoding='utf-8').write(game_cs(km))
    open(sys.argv[2], 'w', encoding='utf-8').write(lobby_cs(km))
    print(sys.argv[1], sys.argv[2])
