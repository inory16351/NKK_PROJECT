"""증강 특수 효과 테이블 생성기 (2026-10-19, 사용자: 칼바람 증강처럼 재밌는 효과 · 높은 등급 중심 · 테이블 따로)
→ Data_Table/증강 특수 효과 테이블.xlsx (Special · Special_Type). 그다음 Tools/xlsx2json.py
증강 테이블에서 effect1 = Special, e1_v1 = 여기 special_id 로 연결. 코드 = Augment/AugmentSpecials.cs (타입별 하나씩).
p1~p5 뜻은 Special_Type 시트. up_param = 강화 레벨마다 커지는 칸 번호 (0 = 없음, 증강 테이블 Aug_Config up_pct 만큼).
증강 설명 글 자리표시 {s1}~{s5} = 이 값 (뒤에 % = ×100)."""
import copy, os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, 'Data_Table', '증강 특수 효과 테이블.xlsx')
STYLE_SRC = os.path.join(ROOT, 'Data_Table', '사람 테이블.xlsx')

# 타입: 칸 뜻 (p1~p5) · 설명
TYPES = [
    ('Giant', '간격 (초)', '마리 수', '크기 배율', '공격력 배율', '지속 (초)', 'p1 초마다 무작위 쥐 p2 마리가 p5 초 동안 거대해짐 (크기 ×p3 · 공격력 ×p4)'),
    ('Snowball', '부술 때마다 공격력 +', '최대', '-', '-', '-', '이번 층에서 물건을 부술 때마다 쥐 공격력 +p1 (최대 p2, 층이 바뀌면 처음부터)'),
    ('Blades', '칼날 수', '반경', '피해 (평균 공격력 ×)', '같은 대상 다시 때리는 간격 (초)', '회전 속도 (바퀴/초)', '쥐 무리 가운데를 도는 치즈 칼날 p1 개: 닿은 물건·고양이·사람·보스에 평균 공격력 × p3'),
    ('Lightning', '간격 (초)', '대상 수', '피해 (평균 공격력 ×)', '-', '-', 'p1 초마다 화면 속 대상 p2 곳에 번개 (평균 공격력 × p3)'),
    ('ClickBomb', '폭발 수', '퍼짐 반경', '피해 (평균 공격력 ×)', '폭발 반경', '-', '클릭 총공격 때 누른 곳 둘레 p2 안에 작은 폭발 p1 번 (반경 p4, 평균 공격력 × p3)'),
    ('Mitosis', '확률', '임시 쥐 시간 (초)', '동시에 최대', '-', '-', '물건을 부수면 p1 확률로 그 자리에 임시 쥐 (p2 초, 최대 p3 마리)'),
    ('Rewind', '되돌리는 시간 (초)', '판당 횟수', '-', '-', '-', '제한시간이 0 이 되면 한 번 p1 초를 되돌림 (판당 p2 번)'),
    ('UltEcho', '확률', '-', '-', '-', '-', '필살기가 끝나면 p1 확률로 다른 쥐의 필살기가 바로 한 번 더'),
    ('Tornado', '간격 (초)', '지속 (초)', '반경', '초당 피해 (평균 공격력 ×)', '이동 속도', 'p1 초마다 쥐 토네이도가 p2 초 동안 열린 방을 휩씀: 반경 p3 안 물건을 끌어당기며 초당 평균 공격력 × p4'),
    ('WallQuake', '다른 벽 체력 감소 (최대 체력 비율)', '-', '-', '-', '-', '벽 하나가 무너지면 남은 벽들이 최대 체력의 p1 만큼 깎임'),
    ('PromoteBoom', '반경', '피해 (평균 공격력 ×)', '-', '-', '-', '승급할 때 새 쥐 자리에서 폭죽 폭발 (반경 p1, 평균 공격력 × p2)'),
    ('Jackpot', '확률', '연구자료 배율', '-', '-', '-', '층 클리어 때 p1 확률로 연구자료 × p2 (잭팟!)'),
]
# id · 타입 · 이름(기획용) · p1~p5 · 강화로 커지는 칸
SPECIALS = [
    (83001, 'Giant', '거대 쥐 출현', 12, 2, 1.8, 3, 6, 4),
    (83002, 'Snowball', '눈덩이 굴리기', 0.01, 0.6, 0, 0, 0, 2),
    (83003, 'Blades', '회전 치즈 칼날', 3, 170, 1.5, 0.4, 0.5, 3),
    (83004, 'Lightning', '천둥 번개', 6, 3, 8, 0, 0, 3),
    (83005, 'ClickBomb', '클릭 폭격', 4, 130, 4, 90, 0, 3),
    (83006, 'Mitosis', '세포 분열', 0.15, 8, 10, 0, 0, 1),
    (83007, 'Rewind', '시간 역행', 30, 1, 0, 0, 0, 1),
    (83008, 'UltEcho', '필살기 메아리', 0.35, 0, 0, 0, 0, 1),
    (83009, 'Tornado', '쥐 토네이도', 15, 6, 200, 3, 260, 4),
    (83010, 'WallQuake', '연쇄 붕괴', 0.15, 0, 0, 0, 0, 1),
    (83011, 'PromoteBoom', '승급 폭죽', 220, 6, 0, 0, 0, 2),
    (83012, 'Jackpot', '잭팟', 0.2, 3, 0, 0, 0, 1),
]


def main():
    wb = openpyxl.load_workbook(STYLE_SRC)
    st = [copy.copy(wb['Human'].cell(r, 1)._style) for r in (1, 2, 3, 4)]
    old = list(wb.sheetnames)

    def sheet(name, head, keys, types, data, widths):
        ws = wb.create_sheet(name)
        for c, (h, k, t) in enumerate(zip(head, keys, types), 1):
            for r, v in enumerate((h, k, t), 1): ws.cell(r, c, v)._style = copy.copy(st[r - 1])
        for i, row in enumerate(data):
            for c, v in enumerate(row, 1): ws.cell(4 + i, c, v)._style = copy.copy(st[3])
        for c, w in enumerate(widths, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w

    sheet('Special', ['특수 효과 id', '타입', '이름 (기획용)', 'p1', 'p2', 'p3', 'p4', 'p5', '강화로 커지는 칸 (0 = 없음)'],
          ['special_id', 'type', '-', 'p1', 'p2', 'p3', 'p4', 'p5', 'up_param'],
          ['int', 'string', '-', 'float', 'float', 'float', 'float', 'float', 'int'],
          [list(s) for s in SPECIALS], [12, 14, 16, 8, 8, 8, 8, 8, 14])
    sheet('Special_Type', ['타입', 'p1', 'p2', 'p3', 'p4', 'p5', '하는 일'], ['-'] * 7, ['-'] * 7,
          [list(t) for t in TYPES], [14, 22, 18, 18, 22, 16, 80])
    for nm in old: del wb[nm]
    wb.save(PATH)
    print('특수 효과', len(SPECIALS), '→', PATH)


if __name__ == '__main__':
    main()
