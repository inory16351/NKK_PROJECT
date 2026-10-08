"""BalanceProbe 측정 결과(RowsCsv) → gen_stage_table.py 의 SPU 표 · 스테이지 테이블 Boss 시트 hp_pow_sec 갱신값 계산.
- 일반 층: 층별 중앙값 % (사용 초 ÷ 기준 초) 가 OK_LO~OK_HI 밖이면 SPU × 측정%/TARGET (한 번에 LO~HI 배 제한).
- 보스 층 (허들, 210초의 90% 목표): 길 뚫기(사용 초 − 보스전) 와 보스전을 따로.
    길 뚫기 = 목표(BOSS_TOTAL_SEC − BOSS_FIGHT_SEC) 대비 비율이 PATH_OK 밖이면 SPU 보정
    보스전   = BOSS_FIGHT_SEC 대비 비율이 FIGHT_OK 밖이면 그 보스 hp_pow_sec 보정
사용: python Tools/update_spu.py <csv> [<csv> ...] [--write]
  --write 면 gen_stage_table.py SPU 와 스테이지 테이블 Boss hp_pow_sec 를 직접 고침 (그다음 gen_stage_table.py · xlsx2json.py)"""
import os, re, sys, csv, statistics
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GEN = os.path.join(ROOT, 'Tools', 'gen_stage_table.py')
XLSX = os.path.join(ROOT, 'Data_Table', '스테이지 테이블.xlsx')
TARGET = 75.0
LO, HI = 0.6, 1.6
OK_LO, OK_HI = 60.0, 90.0       # 일반 층: 이 범위(%)면 괜찮음 → SPU 안 고침 (사용자 2026-10-08: 83~88% 정도는 괜찮음)
PATH_OK = (0.8, 1.2)            # 보스 층 길 뚫기: 목표 대비 이 비율이면 그대로
FIGHT_OK = (0.85, 1.15)         # 보스전: 목표 대비 이 비율이면 그대로
BOSS_EVERY = 5


def load_gen():
    src = open(GEN, encoding='utf-8').read()
    m = re.search(r'SPU = \{[^\n]*\n(.*?)\n\}', src, re.S)
    spu = {int(k): float(v) for k, v in re.findall(r'(\d+):\s*([\d.]+)', m.group(1))}
    num = lambda n: float(re.search(n + r' = ([\d.]+)', src).group(1))
    return src, spu, num('SPU_DEFAULT'), num('BOSS_TOTAL_SEC'), num('BOSS_FIGHT_SEC')


def main():
    files = [a for a in sys.argv[1:] if not a.startswith('--')]
    rows = [r for fn in files for r in csv.DictReader(open(fn, encoding='utf-8'))]
    src, spu, dflt, boss_total, boss_fight = load_gen()
    path_goal = boss_total - boss_fight
    by = {}
    for r in rows: by.setdefault(int(r['floor']), []).append(r)
    new = dict(spu); hp_k = {}
    print('층  회차  %들                중앙값  SPU → 새 SPU')
    allp, inb = [], 0
    for f in sorted(by):
        rs = by[f]; old = spu.get(f, dflt)
        pcts = [100 * float(r['used']) / float(r['base']) for r in rs]; med = statistics.median(pcts)
        if f % BOSS_EVERY == 0:
            fights = [float(r['boss']) for r in rs]
            paths = [float(r['used']) - float(r['boss']) for r in rs]
            pk = statistics.median(paths) / path_goal; fk = statistics.median(fights) / boss_fight
            k = 1.0 if PATH_OK[0] <= pk <= PATH_OK[1] else min(HI, max(LO, pk))
            hk = 1.0 if FIGHT_OK[0] <= fk <= FIGHT_OK[1] else min(2.0, max(0.5, 1 / fk))
            hp_k[f] = hk
            new[f] = round(old * k, 2)
            print(f'{f:>2}B {len(rs)}    {", ".join(f"{p:.0f}" for p in pcts):<18} {med:5.0f}%  {old:6.2f} → {new[f]:6.2f}'
                  f'   길 {statistics.median(paths):.0f}초/{path_goal:.0f} · 보스전 {statistics.median(fights):.0f}초/{boss_fight:.0f} → 체력 ×{hk:.2f}')
            inb += sum(1 for p in pcts if 85 <= p <= 100)
        else:
            k = 1.0 if OK_LO <= med <= OK_HI else min(HI, max(LO, med / TARGET))
            new[f] = round(old * k, 2)
            print(f'{f:>2}  {len(rs)}    {", ".join(f"{p:.0f}" for p in pcts):<18} {med:5.0f}%  {old:6.2f} → {new[f]:6.2f}')
            inb += sum(1 for p in pcts if OK_LO <= p <= OK_HI)
        allp += pcts
    print(f'전체 {len(allp)}번 중 {inb}번이 목표 범위 (일반 {OK_LO:.0f}~{OK_HI:.0f}% · 보스 층 85~100%)')
    if '--write' not in sys.argv: return
    items = sorted(new.items()); lines = []
    for i in range(0, len(items), 10):
        lines.append('    ' + ' '.join(f'{k}: {v},' for k, v in items[i:i + 10]))
    src = re.sub(r'(SPU = \{)[^\n]*\n.*?\n\}', lambda m: m.group(1) + '                               # 측정 BalanceProbe (Tools/update_spu.py 로 갱신)\n' + '\n'.join(lines) + '\n}', src, count=1, flags=re.S)
    open(GEN, 'w', encoding='utf-8').write(src)
    print('gen_stage_table.py SPU 갱신')
    if hp_k:
        wb = openpyxl.load_workbook(XLSX); ws = wb['Boss']
        keys = [c.value for c in ws[2]]; cf, ch = keys.index('floor') + 1, keys.index('hp_pow_sec') + 1
        for r in range(4, ws.max_row + 1):
            f = ws.cell(r, cf).value
            if f in hp_k and hp_k[f] != 1.0:
                v = round(float(ws.cell(r, ch).value) * hp_k[f], 1)
                print(f'  Boss {f}층 hp_pow_sec {ws.cell(r, ch).value} → {v}'); ws.cell(r, ch).value = v
        wb.save(XLSX)


if __name__ == '__main__':
    main()
