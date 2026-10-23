"""증강 테이블 생성기 (2026-10-19) → Data_Table/증강 테이블.xlsx
사용자 결정: 31층 도달 시 해금 → 그 뒤 층 클리어마다 3개 중 1개 · 등급마다 증강 풀이 아예 다름 (롤체·칼바람 증강처럼) ·
 얻으면 지하 '증강 연구소' 보관함에 쌓임 · 판 안 슬롯 2칸에서 시작 · 공용 스킬 Aug_Slot 3개(10 · 12 · 16훈장)로 한 칸씩 최대 5 · 합성(같은 등급 N개 → 상승/유지/하락) ·
 장착한 채 층을 깨면 그 증강 한 개마다 강화 포인트(층수 기준) → 아웃게임에서 포인트 + 치즈로 강화 · 41층 도달 시 1개 들고 가기.
효과 = 공용 스킬 효과 타입 그대로 (CommonSkill 이 공용 스킬 노드와 똑같이 더함) — 공용 스킬 테이블 Common_Effect_Type 시트에 뜻.
설명 글 자리표시: {1} {1b} {1c} = 효과1 밸류_01·02·03, {2} {2b} = 효과2. 뒤에 % 를 붙이면 ×100 ({1%} {1b%}). 강화 레벨만큼 커진 값으로 채움.
고친 뒤: python Tools/gen_augment_table.py && python Tools/gen_string_table.py && python Tools/gen_augment_table.py --en && python Tools/xlsx2json.py"""
import copy, os, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(ROOT, 'Data_Table', '증강 테이블.xlsx')
STYLE_SRC = os.path.join(ROOT, 'Data_Table', '사람 테이블.xlsx')
STR = os.path.join(ROOT, 'Data_Table', '스트링 테이블.xlsx')

# 등급: id · 이름 · 영어 · 색 · 합성 재료 수 · 상승/유지/하락 확률 (사용자: 낮은 등급은 재료 많이·상승 확률 낮게, 높은 등급은 재료 적게·상승 확률 높게, 같은 등급 결과 포함)
GRADES = [
    (1, '브론즈', 'Bronze', '#c98a4b', 4, 0.30, 0.70, 0.00),
    (2, '실버', 'Silver', '#b9c3cc', 3, 0.40, 0.45, 0.15),
    (3, '골드', 'Gold', '#f2c14e', 2, 0.50, 0.35, 0.15),
    (4, '프리즘', 'Prism', '#c58cf0', 2, 0.00, 0.85, 0.15),
]
# 층 클리어 때 나오는 등급 가중치 (그 층 이상부터)
# 2026-10-23 기획서 Ver.04 (증강 32개 = 브론즈 4 · 실버 6 · 골드 12 · 프리즘 10): 브론즈가 4개뿐이라 같은 4개만 반복되지 않게 실버·골드 쪽으로 (수치는 나중에 밸런스 때 다시)
OFFER = [(31, 45, 35, 18, 2), (41, 35, 38, 23, 4), (51, 28, 38, 28, 6), (61, 22, 37, 33, 8), (71, 17, 35, 38, 10), (81, 12, 33, 42, 13), (91, 8, 30, 45, 17)]
# 강화 치즈 (2026-10-19 전체 밸런스): 예전 40 × 치즈 배율 = 층 수입의 0.001% 수준 → 층 치즈 수입 기준으로
UP_CHEESE_FLOORS = 0.15                  # Lv0→1 강화 치즈 = 층 하나 치즈 수입의 이 배 (Lv5 까지 합 ≈ 2층)
try:
    sys_path = __import__('sys').path; sys_path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
    import gen_economy as _ge
    _KA, _KB = _ge.fit_k(); _KB = round(_KB, 4); _UP_CHEESE_BASE = round(_KA * 9 * UP_CHEESE_FLOORS)
except Exception:
    _KB, _UP_CHEESE_BASE = 1.0, 40
CONFIG = [
    ('unlock_floor', 31, '이 층에 도달하면 증강 해금. 이 층 이상을 클리어할 때마다 증강 고르기'),
    ('carry_floor', 41, '이 층에 도달하면 보관함 증강 1개를 들고 판 시작 (증강 연구소에서 고름)'),
    ('slot_base', 2, '판 안 장착 슬롯 기본 수 (공용 스킬 Aug_Slot 합만큼 +, 2026-10-20 사용자: 2칸에서 시작해 조건마다 한 칸씩)'),
    ('slot_max', 5, '판 안 장착 슬롯 최대 (관리 창 자물쇠 미리보기도 여기까지)'),
    ('offer_count', 3, '층 클리어 때 보여 주는 증강 수 (그중 1개)'),
    ('pt_per_floor', 1, '층 클리어 때 장착한 증강 한 개마다 강화 포인트 = 클리어 층 × 이 값'),
    ('up_max', 5, '강화 최대 레벨'),
    ('up_pct', 0.1, '강화 1레벨마다 효과 값 +이 비율 (Lv 5 = ×1.5)'),
    ('up_pt_base', 150, '강화 포인트 필요 = 기본 × 증가^레벨'),
    ('up_pt_grow', 1.6, ''),
    ('up_cheese_base', _UP_CHEESE_BASE, '강화 치즈 = 기본 × 스테이지 테이블 치즈 배율(최고 도달 층) × 층 배율^(min(층, 상한)-1) × 증가^레벨. 기본 = 층 치즈 수입(방 9개)의 UP_CHEESE_FLOORS 배 (Lv0→1)'),
    ('up_cheese_grow', 1.5, ''),
    ('up_cheese_floor_grow', _KB, '층 수입 배율 (Tools/gen_economy.fit_k, 상점 bc_floor_grow 와 같음)'),
    ('up_cheese_floor_cap', 30, '층 수입 배율을 이 층까지만'),
    ('seen_reveal', 1, '1 = 한 번이라도 얻은 증강만 이름·설명 보임 (나머지 ???)'),
]
# 2026-10-23 기획서 Ver.04 로 빠진 증강 → 같은 등급 새 증강 (사용자: 세이브에 있던 것은 자동 교체, 강화 레벨 그대로). Aug_Config 'replace_<옛 id>' = 새 id
REPLACE = {
    81102: 81101, 81103: 81105, 81104: 81106, 81108: 81107, 81109: 81101, 81110: 81105, 81111: 81106, 81112: 81107,
    81201: 81210, 81203: 81207, 81204: 81214, 81205: 81212, 81208: 81206, 81209: 81213, 81211: 81210,
    81301: 81306, 81302: 81315, 81304: 81307, 81305: 81202, 81308: 81313, 81309: 81314, 81310: 81311, 81312: 81303,
    81401: 81403, 81402: 81409, 81404: 81408, 81405: 81413, 81406: 81414, 81407: 81412, 81410: 81415, 81411: 81403,
}
CONFIG += [('replace_%d' % o, n, '없어진 증강 (기획서 Ver.04) → 같은 등급 이 증강으로 (세이브 자동 교체)') for o, n in sorted(REPLACE.items())]

# (id, 등급, 이름, 영어, 설명, 영어 설명, 효과1 (타입, v1, v2, v3), 효과2 또는 None, 아이콘)
# 2026-10-23 기획서 Ver.04 (김규장, 기획서/증강 기획서 Ver.04 반영 결정.md): 수치 증강 10 + 시각 효과 증강 22 = 32개. 수치는 대충 (나중에 밸런스 때 한꺼번에)
A = [
    # ── 브론즈: 수치 ──
    (81101, 1, '날카로운 앞니', 'Sharp Incisors', '쥐 공격력 +{1%}%', 'Rat attack +{1%}%', ('Atk_Pct', 0.12, 0, 0), None, 'aug_tooth'),
    (81105, 1, '강철 앞니', 'Steel Teeth', '벽 피해 +{1%}%', 'Wall damage +{1%}%', ('Wall_Dmg_Pct', 0.2, 0, 0), None, 'aug_wall'),
    (81106, 1, '급소 찾기', 'Weak Spot', '치명타 확률 +{1%}%', 'Critical chance +{1%}%', ('Crit_Chance', 0.04, 0, 0), None, 'aug_crit'),
    (81107, 1, '보스 사냥 면허', 'Boss Hunting License', '보스 피해 +{1%}%', 'Boss damage +{1%}%', ('Boss_Dmg_Pct', 0.2, 0, 0), None, 'aug_boss'),
    # ── 실버: 수치 3 + 한 곳 (클릭 폭격 · 미다스 · 잭팟) ──
    (81210, 2, '프로틴 셰이크', 'Protein Shake', '쥐 공격력 +{1%}%', 'Rat attack +{1%}%', ('Atk_Pct', 0.3, 0, 0), None, 'aug_protein'),
    (81207, 2, '치명적 앞니', 'Deadly Fang', '치명타 확률 +{1%}% · 치명타 피해 +{2%}%', 'Crit chance +{1%}% · Crit damage +{2%}%', ('Crit_Chance', 0.06, 0, 0), ('Crit_Dmg', 0.5, 0, 0), 'aug_fang'),
    (81214, 2, '철거 전문가', 'Demolition Expert', '벽 피해 +{1%}%', 'Wall damage +{1%}%', ('Wall_Dmg_Pct', 0.5, 0, 0), None, 'aug_demolish'),
    (81212, 2, '클릭 폭격', 'Click Barrage', '클릭 총공격 때 누를 때마다 {s1%}% 확률로 폭발 (공격력 ×{s3})', 'Each rush click has a {s1%}% chance to explode (attack ×{s3})', ('Special', 83005, 0, 0), None, 'aug_click'),
    (81206, 2, '미다스의 앞니', 'Midas Bite', '황금 물건을 부수면 치즈 +{s1%}%', 'Smashing golden items gives +{s1%}% cheese', ('Special', 83013, 0, 0), None, 'aug_gold'),
    (81213, 2, '잭팟', 'Jackpot', '층 클리어 때 {s1%}% 확률로 연구자료 ×{s2}', '{s1%}% chance for ×{s2} Research on floor clear', ('Special', 83012, 0, 0), None, 'aug_jackpot'),
    # ── 골드: 수치 2 + 방 일부 (타격 · 쥐 변신 · 방 이벤트 · 등급 통일) ──
    (81306, 3, '보스 킬러', 'Boss Killer', '보스 피해 +{1%}% · 필살기 게이지 +{2%}%', 'Boss damage +{1%}% · Ult charge +{2%}%', ('Boss_Dmg_Pct', 0.6, 0, 0), ('Ult_Gauge_Pct', 0.25, 0, 0), 'aug_skull'),
    (81315, 3, '강철 턱', 'Iron Jaw', '쥐 공격력 +{1%}% · 벽 피해 +{2%}%', 'Rat attack +{1%}% · Wall damage +{2%}%', ('Atk_Pct', 0.5, 0, 0), ('Wall_Dmg_Pct', 0.4, 0, 0), 'aug_jaw'),
    (81307, 3, '다다다다', 'Rat-a-tat', '부딪힐 때 {1%}% 확률로 {1b}번 더 공격', '{1%}% chance to hit {1b} more times', ('Multi_Hit', 0.25, 3, 0), None, 'aug_rapid'),
    (81202, 3, '찌릿찌릿 이빨', 'Zappy Teeth', '물 때 {1%}% 확률로 주변 물건 {1b}개 감전', '{1%}% chance on bite to zap {1b} nearby items', ('Bite_Zap', 0.08, 3, 0.6), None, 'aug_zap'),
    (81313, 3, '연쇄 붕괴', 'Domino Walls', '벽 하나가 무너지면 남은 벽 체력 -{s1%}%', 'When a wall falls, other walls lose {s1%}% HP', ('Special', 83010, 0, 0), None, 'aug_domino'),
    (81314, 3, '승급 폭죽', 'Promotion Fireworks', '승급할 때 새 쥐 자리에서 폭죽 폭발 (공격력 ×{s2})', 'Promotions set off fireworks at the new rat (attack ×{s2})', ('Special', 83011, 0, 0), None, 'aug_rankboom'),
    (81316, 3, '쥐 탑 쌓기', 'Rat Tower', '장착하면 쥐 {s2}마리씩 쌓아 합친 쥐 탑 {s1}개 (능력치 합산, 판 끝까지)', 'Stack {s2} rats into a tower, up to {s1} towers (stats combined, lasts the run)', ('Special', 83014, 0, 0), None, 'aug_tower'),
    (81311, 3, '천둥 번개', 'Thunderstorm', '{s1}초마다 먹구름에서 번개 {s2}개 (공격력 ×{s3})', 'Every {s1}s, {s2} bolts fall from the storm cloud (attack ×{s3})', ('Special', 83004, 0, 0), None, 'aug_storm'),
    (81317, 3, '치즈 비', 'Cheese Rain', '{s1}초마다 {s2}초 동안 치즈 비 (치즈 +{s3%}%)', 'Every {s1}s, cheese rains for {s2}s (cheese +{s3%}%)', ('Special', 83015, 0, 0), None, 'aug_cheeserain'),
    (81303, 3, '탄생 축제', 'Birth Party', '새끼가 태어나면 주변 쥐 {1}초 광란', 'Births send nearby rats into a {1}s frenzy', ('Birth_Frenzy', 4, 250, 1.6), None, 'aug_party'),
    (81318, 3, '골든 햄스터 생쥐', 'Golden Hamster Rats', '일반 쥐 전원 골든 햄스터 · 늘 공으로 굴러다님 (이후 태어나는 쥐도)', 'All Common rats become Golden Hamsters, always rolling', ('Special', 83016, 0, 0), None, 'aug_u_hamster'),
    (81319, 3, '판다 잠옷 정글리안 생쥐', 'Panda Pajama Rats', '일반 쥐 전원 판다 잠옷 정글리안 · 늘 공으로 굴러다님 (이후 태어나는 쥐도)', 'All Common rats become Panda Pajama hamsters, always rolling', ('Special', 83017, 0, 0), None, 'aug_u_panda'),
    # ── 프리즘: 수치 1 + 방 전체 ──
    (81403, 4, '전설의 앞니', 'Legendary Incisors', '쥐 공격력 +{1%}% · 모든 피해 +{2%}%', 'Rat attack +{1%}% · All damage +{2%}%', ('Atk_Pct', 1.0, 0, 0), ('Dmg_Pct', 0.3, 0, 0), 'aug_crown'),
    (81409, 4, '회전 치즈 칼날', 'Spinning Cheese Blades', '칠 때 {s1%}% 확률로 그 쥐 둘레를 치즈 칼날 {s2}개가 {s5}초 돌며 공격 (공격력 ×{s4})', '{s1%}% on hit: {s2} cheese blades spin around the rat for {s5}s (attack ×{s4})', ('Special', 83003, 0, 0), None, 'aug_blades'),
    (81408, 4, '거대 쥐 출현', 'Rise of the Giants', '{s1}초마다 쥐 {s2}마리가 {s5}초 동안 거대화 (공격력 ×{s4})', 'Every {s1}s, {s2} rats turn giant for {s5}s (attack ×{s4})', ('Special', 83001, 0, 0), None, 'aug_giant'),
    (81413, 4, '분신 소환', 'Clone Summon', '{s1}초마다 가장 강한 쥐의 거대 분신 {s2}마리가 {s3}초 동안 싸움', 'Every {s1}s, {s2} giant clones of your strongest rat fight for {s3}s', ('Special', 83018, 0, 0), None, 'aug_clones'),
    (81414, 4, '운석 소환', 'Meteor Strike', '{s1}초마다 큰 치즈 운석 낙하 (공격력 ×{s3})', 'Every {s1}s a giant cheese meteor falls (attack ×{s3})', ('Special', 83019, 0, 0), None, 'aug_meteor'),
    (81412, 4, '쥐 토네이도', 'Rat Tornado', '{s1}초마다 토네이도가 {s2}초 동안 방을 휩씀 (초당 공격력 ×{s4})', 'Every {s1}s a tornado sweeps the room for {s2}s (attack ×{s4}/s)', ('Special', 83009, 0, 0), None, 'aug_tornado'),
    (81415, 4, '장난감 기차', 'Toy Train', '{s1}초마다 기차가 맵을 가로지름 (공격력 ×{s3})', 'Every {s1}s a toy train crosses the map (attack ×{s3})', ('Special', 83020, 0, 0), None, 'aug_train'),
    (81416, 4, '스케이터 생쥐', 'Skater Rats', '레어 쥐 전원 양복 대머리 생쥐 · 매우 빠르게 달리며 박치기 (이후 태어나는 쥐도)', 'All Rare rats become suited bald rats that dash and headbutt', ('Special', 83021, 0, 0), None, 'aug_u_skater'),
    (81417, 4, '맘스땃쥐 생쥐', 'Burger Rats', '에픽 쥐 전원 맘스땃쥐 · 패티 최대 {s5}장 · 박치기 피해 +{s4%}% (이후 태어나는 쥐도)', 'All Epic rats become Burger rats · up to {s5} patties · headbutt +{s4%}%', ('Special', 83022, 0, 0), None, 'aug_u_burger'),
    (81418, 4, '5G 생쥐', '5G Rats', '에픽 쥐 전원 5G · 동료 수와 상관없이 모든 쥐에게 데이터 전송 (이후 태어나는 쥐도)', 'All Epic rats become 5G rats that boost every rat', ('Special', 83023, 0, 0), None, 'aug_u_fiveg'),
]
# 강화로 커지지 않는 칸 (개수·간격·켜기 같은 값): (효과 타입, 밸류 번호)
NO_SCALE = {('Multi_Hit', 2), ('Bite_Zap', 2), ('Birth_Frenzy', 2), ('Birth_Frenzy', 3)}


def style_of():
    # 스타일은 같은 통합 문서 안에서만 복사됨 → 사람 테이블을 열어 그 안에 새 시트를 만들고 원래 시트는 지움
    wb = openpyxl.load_workbook(STYLE_SRC)
    return wb, [copy.copy(wb['Human'].cell(r, 1)._style) for r in (1, 2, 3, 4)]


def sheet(wb, st, name, head, keys, types, data, widths):
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name)
    for c, (h, k, t) in enumerate(zip(head, keys, types), 1):
        for r, v in enumerate((h, k, t), 1): ws.cell(r, c, v)._style = copy.copy(st[r - 1])
    for i, row in enumerate(data):
        for c, v in enumerate(row, 1): ws.cell(4 + i, c, v)._style = copy.copy(st[3])
    for c, w in enumerate(widths, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    return ws


def build():
    wb, st = style_of()
    old = list(wb.sheetnames)
    rows = []
    for (i, g, nm, en, d, den, e1, e2, icon) in A:
        e2 = e2 or ('None', 0, 0, 0)
        rows.append([i, g, nm, d, e1[0], e1[1], e1[2], e1[3], e2[0], e2[1], e2[2], e2[3], icon, 1, en])
    sheet(wb, st, 'Augment',
          ['증강 id', '등급', '이름', '설명 (자리표시 {1} {1b} {2} …, % = ×100)', '효과1', '효과1 밸류_01', '효과1 밸류_02', '효과1 밸류_03', '효과2', '효과2 밸류_01', '효과2 밸류_02', '효과2 밸류_03', '아이콘 (Art/Augment)', '나옴 (0 = 안 나옴)', '영어 (번역 참고)'],
          ['augment_id', 'grade', 'aug_name', 'aug_desc', 'effect1', 'e1_v1', 'e1_v2', 'e1_v3', 'effect2', 'e2_v1', 'e2_v2', 'e2_v3', 'icon', 'enabled', '-'],
          ['int', 'int', 'string', 'string', 'string', 'float', 'float', 'float', 'string', 'float', 'float', 'float', 'string', 'int', '-'],
          rows, [10, 6, 18, 40, 16, 10, 10, 10, 16, 10, 10, 10, 16, 8, 24])
    sheet(wb, st, 'Grade',
          ['등급', '이름', '색', '합성 재료 수 (같은 등급)', '상승 확률', '유지 확률', '하락 확률', '영어'],
          ['grade', 'grade_name', 'color', 'synth_need', 'synth_up', 'synth_same', 'synth_down', '-'],
          ['int', 'string', 'string', 'int', 'float', 'float', 'float', '-'],
          [[g, n, c, need, u, s, d, en] for g, n, en, c, need, u, s, d in GRADES], [6, 10, 10, 16, 10, 10, 10, 10])
    sheet(wb, st, 'Offer',
          ['이 층부터 (클리어 층)', '브론즈 가중치', '실버', '골드', '프리즘'],
          ['from_floor', 'w1', 'w2', 'w3', 'w4'], ['int', 'float', 'float', 'float', 'float'],
          [list(o) for o in OFFER], [16, 12, 10, 10, 10])
    sheet(wb, st, 'Aug_Config', ['키', '값', '설명'], ['key', 'value', '-'], ['string', 'float', '-'],
          [[k, v, d] for k, v, d in CONFIG], [18, 10, 70])
    sheet(wb, st, 'No_Scale', ['효과 타입', '밸류 번호', '설명'], ['effect', 'index', '-'], ['string', 'int', '-'],
          [[e, k, '강화해도 안 커지는 칸 (개수·간격·켜기)'] for e, k in sorted(NO_SCALE)], [18, 10, 40])
    ws = sheet(wb, st, 'Column_Desc', ['칸', '설명'], ['-', '-'], ['-', '-'],
               [['생성', 'Tools/gen_augment_table.py 로 생성 (내용은 파이썬 A 목록을 고치고 다시 실행)'],
                ['Special', 'effect1 = Special 이면 e1_v1 = 증강 특수 효과 테이블 special_id (칼바람식 특수 효과, 설명 자리표시 {s1}~{s5})'],
                ['effect1 / effect2', '공용 스킬 효과 타입 (공용 스킬 테이블 Common_Effect_Type 시트). 판 안에서 장착한 증강은 공용 스킬 노드처럼 더해짐. Research_Pct = 층 클리어 연구자료 ×(1 + 합)'],
                ['Grade.synth_*', '같은 등급 synth_need 개 → 무작위 증강 1개, 그 등급에서 상승/유지/하락. 맨 아래 등급의 하락 · 맨 위 등급의 상승은 유지로'],
                ['Offer', '층 클리어 때 3장 등급 가중치 (클리어 층 기준, from_floor 이하 중 가장 높은 줄)']], [20, 100])
    for nm in old: del wb[nm]
    wb.save(PATH)
    print('증강', len(A), '종 →', PATH)


def fill_en():
    # gen_string_table 뒤: 증강 이름·설명·등급 이름의 영어 칸 채우기 (번역 필요 → 비움)
    aw = openpyxl.load_workbook(PATH, data_only=True)
    en = {}
    for ws, kcol, ecol in ((aw['Augment'], 'aug_name', None), (aw['Augment'], 'aug_desc', None), (aw['Grade'], 'grade_name', None)):
        K = [c.value for c in ws[2]]
        for r in ws.iter_rows(min_row=4, values_only=True):
            if r[0] is None: continue
            sk = r[K.index(kcol + '_sk')] if kcol + '_sk' in K else None
            if not sk: continue
            if ws.title == 'Augment':
                row = next(a for a in A if a[0] == r[0])
                en[sk] = row[3] if kcol == 'aug_name' else row[5]
            else:
                en[sk] = next(g[2] for g in GRADES if g[0] == r[0])
    sw = openpyxl.load_workbook(STR); ws = sw['String']; K = [c.value for c in ws[2]]
    n = 0
    for r in ws.iter_rows(min_row=4):
        k = r[0].value
        if k in en and r[K.index('en')].value != en[k]:
            r[K.index('en')].value = en[k]; r[K.index('status')].value = None; n += 1
    sw.save(STR)
    print('영어 채움', n)


if __name__ == '__main__':
    if '--en' in sys.argv: fill_en()
    else: build()
