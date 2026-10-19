"""스트링 테이블 생성기 → Data_Table/스트링 테이블.xlsx (그다음 Tools/xlsx2json.py)

다른 테이블의 화면에 나오는 글 칸(아래 COLS) 옆에 `<칸>_sk` 키 칸(int)을 두고, 그 키로 스트링 테이블(String 시트)의 ko·en 을 찾음.
- 다른 테이블의 한글은 그대로 둠 (사람이 읽는 용). 키는 관리·코드용.
- 같은 칸에서 같은 한글이면 같은 키 (공용 스킬처럼 이름이 반복되는 칸).
- 키 = 칸 번호(prefix) × 100000 + 행 id (id 가 int 이고 비어 있으면), 아니면 그 칸 범위에서 빈 번호.
- 한글을 고치면: 그 키를 혼자 쓰면 ko 를 고치고 상태 '한글 바뀜' (en 다시 확인), 여럿이 같이 쓰면 새 키.
- 시트를 통째로 다시 만드는 생성기(gen_skill_tree.py 등) 뒤에 돌리면 키 칸이 사라져도 같은 한글이면 같은 키로 다시 채움.
- 씬 글(80xxxxx)·코드 글(90xxxxx) 줄은 이 스크립트가 건드리지 않음 (스트링 테이블에서 직접 관리).
순서: gen_stage_table → gen_economy → gen_skill_tree → gen_string_table → xlsx2json
사용: python Tools/gen_string_table.py"""
import copy, os, re
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'Data_Table')
OUT = os.path.join(SRC, '스트링 테이블.xlsx')
HANGUL = re.compile('[가-힣]')
STEP = 100000

# (엑셀 파일, 시트, id 칸 (없으면 None), 글 칸, prefix) — 화면에 나오는 글만 (기획 메모 desc 등은 뺌)
COLS = [
    ('쥐 캐릭터 테이블', 'Character', 'character_id', 'character_name', 10),
    ('쥐 캐릭터 테이블', 'Character', 'character_id', 'character_desc', 11),
    ('쥐 캐릭터 테이블', 'Skill', 'skill_id', 'skill_name', 12),
    ('쥐 캐릭터 테이블', 'Skill', 'skill_id', 'skill_explain', 13),
    ('쥐 캐릭터 테이블', 'Ultimate', 'ultimate_id', 'ultimate_name', 14),
    ('쥐 캐릭터 테이블', 'Ultimate', 'ultimate_id', 'dev_desc', 15),
    ('쥐 캐릭터 테이블', 'Ultimate', 'ultimate_id', 'ult_line', 16),
    ('쥐 캐릭터 테이블', 'Ult_Caption', None, 'text', 17),
    ('쥐 등급 테이블', 'Grade', None, 'grade_name', 18),
    ('쥐 성장 테이블', 'Growth_Node', 'node_id', 'node_name', 19),
    ('쥐 성장 테이블', 'Growth_Node', 'node_id', 'node_desc', 20),
    ('고양이 캐릭터 테이블', 'Character', 'character_id', 'character_name', 21),
    ('고양이 캐릭터 테이블', 'Skill', 'skill_id', 'skill_name', 22),
    ('공용 스킬 테이블', 'Common_Skill', 'skill_id', 'skill_name', 24),
    ('공용 스킬 테이블', 'Common_Skill', 'skill_id', 'skill_explain', 25),
    ('공용 스킬 테이블', 'Branch', None, 'branch_name', 26),
    ('물건 테이블', 'Item', 'item_id', 'item_name', 27),
    ('물건 테이블', 'Zone', 'zone_id', 'zone_name', 28),
    ('보스 테이블', 'Boss', 'boss_id', 'boss_name', 29),
    ('보스 테이블', 'Boss_Line', 'line_id', 'text', 30),
    ('보스 테이블', 'Atk_Type', None, 'atk_name', 31),
    ('보스 테이블', 'Atk_Type', None, 'line', 32),
    ('사람 테이블', 'Human', 'human_id', 'human_name', 33),
    ('사람 테이블', 'Human_Line', 'line_id', 'text', 34),
    ('상점 테이블', 'Shop_Goods', 'goods_id', 'goods_name', 35),
    ('업적 테이블', 'Achievement', 'achv_id', 'achv_name', 36),
    ('업적 테이블', 'Achievement', 'achv_id', 'achv_desc', 37),
    ('튜토리얼 테이블', 'Speaker', None, 'speaker_name', 38),
    ('튜토리얼 테이블', 'Line', 'line_id', 'text', 39),
    ('튜토리얼 테이블', 'Story', 'line_id', 'text', 40),
    ('티어 테이블', 'Tier', 'tier', 'tier_name', 41),
    ('쥐 캐릭터 테이블', 'Super_Ult', 'super_id', 'super_name', 42),
    ('쥐 캐릭터 테이블', 'Super_Ult', 'super_id', 'super_line', 43),
    ('쥐 캐릭터 테이블', 'Super_Ult', 'super_id', 'finish_line', 44),
    ('보스 테이블', 'Atk_Type', None, 'popup', 45),          # 밈 기술 팝업 (2026-10-19, 100층 확장)
    ('증강 테이블', 'Augment', 'augment_id', 'aug_name', 46),   # 증강 (2026-10-19)
    ('증강 테이블', 'Augment', 'augment_id', 'aug_desc', 47),
    ('증강 테이블', 'Grade', None, 'grade_name', 48),
]

HEAD = (['스트링 키', '한국어', '영어', '출처', '상태', '번역 메모'],
        ['string_key', 'ko', 'en', 'source', 'status', 'note'],
        ['int', 'string', 'string', '-', '-', '-'])
GLOSSARY_HEAD = (['한국어', '영어', '메모'], ['ko', 'en', 'note'], ['-', '-', '-'])


def load_strings():
    rows = {}
    if not os.path.exists(OUT): return rows, None
    wb = openpyxl.load_workbook(OUT)
    ws = wb['String']
    keys = [c.value for c in ws[2]]
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] is None: continue
        d = dict(zip(keys, r)); d['string_key'] = int(d['string_key'])
        rows[d['string_key']] = {k: ('' if v is None else v) for k, v in d.items()}
    return rows, wb


def ensure_sk(ws, col):
    """글 칸 바로 오른쪽에 <col>_sk 칸이 없으면 끼워 넣음. (글 칸 번호, 키 칸 번호) — 1부터"""
    keys = [c.value for c in ws[2]]
    ci = keys.index(col) + 1
    sk = col + '_sk'
    if ci < len(keys) and keys[ci] == sk: return ci, ci + 1
    if sk in keys:                                        # 다른 자리에 있으면 그대로 씀
        return ci, keys.index(sk) + 1
    ws.insert_cols(ci + 1)
    for r, v in ((1, f'{ws.cell(1, ci).value or col} 키'), (2, sk), (3, 'int')):
        c = ws.cell(r, ci + 1, v); c._style = copy.copy(ws.cell(r, ci)._style)
    for r in range(4, ws.max_row + 1): ws.cell(r, ci + 1)._style = copy.copy(ws.cell(r, ci)._style)
    ws.column_dimensions[openpyxl.utils.get_column_letter(ci + 1)].width = 10
    return ci, ci + 1


def main():
    strings, _ = load_strings()
    books = {}
    cells = []          # (prefix, src, ws, row, key칸, 한글, id, 기존 키)
    for f, sh, idc, col, pre in COLS:
        if f not in books: books[f] = openpyxl.load_workbook(os.path.join(SRC, f + '.xlsx'))
        ws = books[f][sh]
        ci, ki = ensure_sk(ws, col)
        keys = [c.value for c in ws[2]]
        ii = keys.index(idc) + 1 if idc else None
        src = f'{f} · {sh} · {col}'
        for r in range(4, ws.max_row + 1):
            if ws.cell(r, 1).value is None: continue
            v = ws.cell(r, ci).value
            if not isinstance(v, str) or v.strip() in ('', '-'):          # 한글이 없는 이름(5G 등)도 번역 대상
                ws.cell(r, ki).value = None; continue
            rid = ws.cell(r, ii).value if ii else None
            old = ws.cell(r, ki).value
            cells.append(dict(pre=pre, src=src, ws=ws, row=r, ki=ki, ko=v, id=rid if isinstance(rid, int) else None,
                              label=str(ws.cell(r, ii).value if ii else ws.cell(r, 1).value), old=int(old) if isinstance(old, (int, float)) else None))

    # 1) 기존 키가 맞으면 (같은 칸 범위 + 같은 한글) 그대로
    used = {}
    by_src_ko = {}
    for k, s in strings.items():
        by_src_ko.setdefault((k // STEP, s['ko']), k)
    pending = []
    for c in cells:
        k = c['old']
        if k and k // STEP == c['pre'] and k in strings and strings[k]['ko'] == c['ko']:
            c['key'] = k
        elif (c['pre'], c['ko']) in by_src_ko:
            c['key'] = by_src_ko[(c['pre'], c['ko'])]
        else:
            pending.append(c)
            continue
        used.setdefault(c['key'], []).append(c)

    # 2) 한글이 바뀐 칸: 그 키를 아무도 안 쓰면 ko 고침, 아니면 새 키
    changed = set()
    fresh = []
    for c in pending:
        k = c['old']
        if k and k // STEP == c['pre'] and k in strings and k not in used:
            strings[k]['ko'] = c['ko']; c['key'] = k; changed.add(k)
            used.setdefault(k, []).append(c); by_src_ko[(c['pre'], c['ko'])] = k
        else:
            fresh.append(c)

    # 3) 새 키
    def free(pre, want):
        if want is not None and want not in strings: return want
        k = pre * STEP + 1
        while k in strings: k += 1
        return k
    for c in fresh:
        if (c['pre'], c['ko']) in by_src_ko:
            c['key'] = by_src_ko[(c['pre'], c['ko'])]
        else:
            k = free(c['pre'], c['pre'] * STEP + c['id'] % STEP if c['id'] is not None else None)
            strings[k] = dict(string_key=k, ko=c['ko'], en='', source='', status='', note='')
            c['key'] = k; by_src_ko[(c['pre'], c['ko'])] = k
        used.setdefault(c['key'], []).append(c)

    for c in cells: c['ws'].cell(c['row'], c['ki']).value = c['key']

    # 출처 · 상태
    table_pres = {p for *_, p in COLS}
    for k, s in strings.items():
        if k // STEP not in table_pres:
            if not s['en']: s['status'] = '번역 필요'
            elif s['status'] == '번역 필요': s['status'] = ''
            continue
        cs = used.get(k)
        if not cs:
            s['status'] = '안 쓰임'; continue
        ids = [c['label'] for c in cs]
        s['source'] = f"{cs[0]['src']} · {', '.join(ids[:6])}{' 외 %d' % (len(ids) - 6) if len(ids) > 6 else ''}"
        if k in changed and s['en']: s['status'] = '한글 바뀜'
        elif not s['en']: s['status'] = '번역 필요'
        elif s['status'] in ('번역 필요', '안 쓰임'): s['status'] = ''

    for f, wb in books.items(): wb.save(os.path.join(SRC, f + '.xlsx'))
    write(strings)
    n = len(strings); todo = sum(1 for s in strings.values() if s['status'] == '번역 필요')
    print(f'스트링 {n}개 (번역 필요 {todo} · 한글 바뀜 {sum(1 for s in strings.values() if s["status"] == "한글 바뀜")} · 안 쓰임 {sum(1 for s in strings.values() if s["status"] == "안 쓰임")})')


def write(strings):
    old = openpyxl.load_workbook(OUT) if os.path.exists(OUT) else None
    style_src = openpyxl.load_workbook(os.path.join(SRC, '사람 테이블.xlsx'))['Human']
    st = [copy.copy(style_src.cell(r, 1)._style) for r in (1, 2, 3, 4)]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'String'
    for c in range(len(HEAD[0])):
        for r in range(3): ws.cell(r + 1, c + 1, HEAD[r][c])._style = copy.copy(st[r])
    for i, k in enumerate(sorted(strings), 4):
        s = strings[k]
        for c, key in enumerate(HEAD[1], 1):
            cell = ws.cell(i, c, s.get(key) or None); cell._style = copy.copy(st[3])
            cell.alignment = openpyxl.styles.Alignment(wrap_text=key in ('ko', 'en'), vertical='top')
    for c, w in zip('ABCDEF', (11, 50, 50, 44, 11, 30)): ws.column_dimensions[c].width = w
    ws.freeze_panes = 'B4'
    # 용어집 (번역 기준 · 사람이 보는 용)
    g = wb.create_sheet('Glossary')
    if old and 'Glossary' in old.sheetnames:
        for row in old['Glossary'].iter_rows(values_only=True): g.append(list(row))
    else:
        for r in range(3): g.append(list(GLOSSARY_HEAD[r]))
    for r in range(1, 4):
        for c in range(1, 4): g.cell(r, c)._style = copy.copy(st[r - 1])
    for c, w in zip('ABC', (24, 30, 60)): g.column_dimensions[c].width = w
    wb.save(OUT)


if __name__ == '__main__':
    main()
