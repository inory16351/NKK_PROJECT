"""엑셀 밸런스 시뮬레이터 (2026-10-19, 사용자: 유니티 안에서 돌리면 토큰이 많이 드니 엑셀로) → Data_Table/밸런스 시뮬레이터.xlsx
유니티 없이 지금 테이블 JSON(Assets/Data) + Tools/balance_sim.py 모델로 처음~100층 진행을 어림잡음.

시트
 · 설정   : 엑셀에서 직접 고치는 입력 (노란 칸). 시간 모델 · 훈장별 끝/시작 전투력 · 훈장별 추가 비용 배율 · 증강 슬롯.
            다시 실행해도 이 시트 값은 유지됨 (지우면 기본값으로 다시 만듦).
 · 층별   : 훈장 t 의 '끝 상태(트리 전부)' 로 각 층이 몇 초 걸리는지 — 엑셀 수식이라 설정을 고치면 바로 바뀜.
            빨강 = 제한시간 넘음(실패) · 주황 = 제한시간 85% 이상(아슬아슬) · 초록 = 40% 미만(쉬움)
 · 요약 · 판_클릭 · 판_미클릭 : 이 스크립트를 실행하면 채움 (판마다 무엇을 살지는 반복 계산이라 파이썬이 계산).
사용: 엑셀에서 설정 고치고 저장 → python Tools/balance_sim_xlsx.py → 엑셀 다시 열기
테이블 곡선 자체(적정·비용)를 바꾸려면 생성기(gen_stage_table · gen_economy · gen_expansion_data · gen_skill_tree)를 고쳐 다시 생성."""
import collections, json, os, sys
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import balance_sim as b

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'Data_Table', '밸런스 시뮬레이터.xlsx')
HURDLES = [5, 10, 15, 16, 19, 30, 40, 50, 60, 70, 80, 90, 100]
MARKS = [16, 30, 31, 41, 50, 70, 100]          # 이 층 '클리어' 시각 (요약)
YELLOW = PatternFill('solid', fgColor='FFF2CC'); HEAD = PatternFill('solid', fgColor='D9E1F2'); BOLD = Font(bold=True)

# 설정 시트 칸 (행 고정 — 층별 수식이 이 주소를 씀)
PARAMS = [   # (이름, 키, 기본값, 설명)
    ('기본 제한시간 (초)', 'BASE_TIME', b.BASE_TIME, 'RunTimer.baseTime'),
    ('시간 배율 C', 'TIME_C', b.TIME_C, '층 시간 = 목표 초 × C × r^-지수 (r = 전투력 ÷ 적정)'),
    ('r<1 지수', 'TIME_LO', b.TIME_LO, '전투력이 모자랄 때 느려지는 정도'),
    ('r≥1 지수', 'TIME_HI', b.TIME_HI, '전투력이 넘칠 때 빨라지는 정도'),
    ('클릭 거의 안 함: 층 시간 배율', 'PASSIVE_SLOW', b.PASSIVE_SLOW, '같은 상태에서 몇 배 느린지 (측정 필요 — 지금은 가정)'),
    ('층 사이 연출 (초)', 'FLOOR_OVERHEAD', b.FLOOR_OVERHEAD, ''),
    ('판마다 로비 시간 (초)', 'RUN_OVERHEAD', b.RUN_OVERHEAD, ''),
    ('판 첫 층 손해 (초)', 'RAMP_SEC', b.RAMP_SEC, '쥐가 불어나는 동안'),
    ('16층~ 훈장·트리에 쓰는 수입 몫', 'LATE_SHARE', b.LATE_SHARE, '나머지 = 뽑기·상점·증강 강화'),
    ('층당 출생 ÷ 최대 마릿수', 'BIRTH_PER_POP', b.BIRTH_PER_POP, '조각 = 태어난 종 +1 (조각 합·도감 조건). 레벨업 조각 = 쥐 등급 테이블 shard_k'),
    ('증강 1칸당 전투력 +', 'AUG_PER_SLOT', b.AUG_PER_SLOT, '31층~ (gen_stage_table 과 같게)'),
    ('시뮬 최대 시간 (h)', 'HOURS_CAP', 60, '못 깨면 여기서 멈춤'),
    ('목표: 클릭 적극 (h)', 'GOAL_ACTIVE', 10, '사용자 2026-10-19'),
    ('목표: 거의 안 함 (h)', 'GOAL_PASSIVE', 15, ''),
]
P_ROW0 = 4                                     # 첫 설정 행
T_ROW0 = P_ROW0 + len(PARAMS) + 3              # 훈장 표 머리 행
T_COLS = ['훈장', '시작 전투력', '끝 전투력', '추가 비용 배율', '증강 슬롯', '제한시간 + (누적)', '보스 층 + (누적)', '승급 연구자료 (테이블)', '조건 층', '최대 마릿수']
LAYOUT = 'v2'                                   # 설정 시트 칸 배치 (바뀌면 설정 시트를 새로 만듦)


def tier_static():
    out = {}
    for t in range(1, 17):
        ta = sum(s['value_01'] for s in b.SKILL if s['tier'] <= t and s['effect_type'] == 'Time_Add')
        bt = sum(s['value_02'] for s in b.SKILL if s['tier'] <= t and s['effect_type'] == 'Boss_Dmg_Pct')
        sl = 2 + sum(s['value_01'] for s in b.SKILL if s['tier'] <= t and s['effect_type'] == 'Aug_Slot')
        tr = b.TIER[t]; fl = next((tr[f'cond{i}_value'] for i in range(1, 5) if tr.get(f'cond{i}_type') in ('Max_Floor', 'Boss_Clear')), 0)
        out[t] = dict(ta=ta, bt=bt, slots=sl, cost=tr['research_cost'], floor=fl)
    return out


def make_settings(wb, st):
    ws = wb.create_sheet('설정', 0)
    ws['A1'] = '밸런스 시뮬레이터 설정 — 노란 칸을 고치고 저장한 뒤 python Tools/balance_sim_xlsx.py'; ws['A1'].font = Font(bold=True, size=13)
    ws['Z1'] = LAYOUT
    ws['A2'] = '층별 시트는 엑셀 수식이라 바로 바뀜 · 요약/판 시트는 스크립트를 다시 돌려야 바뀜'
    for c, h in enumerate(('항목', '키', '값', '설명'), 1):
        x = ws.cell(P_ROW0 - 1, c, h); x.fill = HEAD; x.font = BOLD
    for i, (n, k, v, d) in enumerate(PARAMS):
        r = P_ROW0 + i
        ws.cell(r, 1, n); ws.cell(r, 2, k); ws.cell(r, 3, v).fill = YELLOW; ws.cell(r, 4, d)
    ws.cell(T_ROW0 - 1, 1, '훈장별 (전투력 = BalanceProbe 측정 Tools/probe_results/power_by_tier.json · 추가 비용 배율 = 테이블 비용에 더 곱함, 1 = 그대로)').font = BOLD
    for c, h in enumerate(T_COLS, 1):
        x = ws.cell(T_ROW0, c, h); x.fill = HEAD; x.font = BOLD
    pw = json.load(open(b.POWER, encoding='utf-8'))
    for t in range(1, 17):
        r = T_ROW0 + t
        vals = [t, pw[str(t)]['start'], pw[str(t)]['end'], 1.0, st[t]['slots'], st[t]['ta'], st[t]['bt'], st[t]['cost'], st[t]['floor'], pw[str(t)].get('pop', 100)]
        for c, v in enumerate(vals, 1):
            x = ws.cell(r, c, v)
            if c in (2, 3, 4, 5, 10): x.fill = YELLOW
            if c in (2, 3, 8): x.number_format = '#,##0'
    for c, w in enumerate((30, 16, 14, 50, 10, 16, 16, 20, 10), 1): ws.column_dimensions[L(c)].width = w
    return ws


def read_settings(ws):
    p = {ws.cell(P_ROW0 + i, 2).value: ws.cell(P_ROW0 + i, 3).value for i in range(len(PARAMS))}
    tiers = {}
    for t in range(1, 17):
        r = T_ROW0 + t
        tiers[t] = dict(start=float(ws.cell(r, 2).value), end=float(ws.cell(r, 3).value), cost=float(ws.cell(r, 4).value or 1), slots=float(ws.cell(r, 5).value or 0), pop=float(ws.cell(r, 10).value or 100))
    return p, tiers


def addr(key):                                  # 설정 시트 값 칸 절대 주소
    i = [k for _, k, _, _ in PARAMS].index(key)
    return f"'설정'!$C${P_ROW0 + i}"


def floor_sheet(wb, st):
    if '층별' in wb.sheetnames: del wb['층별']
    ws = wb.create_sheet('층별')
    ws['A1'] = '훈장 끝 상태(트리 전부)로 이 층을 깨는 데 걸리는 초 (수식) — 빨강 = 실패 · 주황 = 제한시간 85%↑ · 초록 = 40%↓ · 굵은 행 = 허들 층'
    heads = ['층', '유형', '보스', '적정 전투력', '목표 초', '연구자료', '시작 층 상한'] + [f'{t}훈장' for t in range(1, 17)]
    for c, h in enumerate(heads, 1):
        x = ws.cell(2, c, h); x.fill = HEAD; x.font = BOLD; x.alignment = Alignment(horizontal='center')
    C, LO, HI, BT = addr('TIME_C'), addr('TIME_LO'), addr('TIME_HI'), addr('BASE_TIME')
    AUG = addr('AUG_PER_SLOT')
    for f in range(1, b.LAST + 1):
        r = 2 + f; s = b.STAGE[f]; boss = f in b.BOSS_FLOORS
        for c, v in enumerate([f, s['floor_type'], '보스' if boss else '', s['pow_need'], s['target_sec'], s['research'], s['skip_cap']], 1):
            x = ws.cell(r, c, v)
            if c in (4, 6): x.number_format = '#,##0'
            if f in HURDLES: x.font = BOLD
        for t in range(1, 17):
            tr = T_ROW0 + t; col = 7 + t
            R = f"('설정'!$C${tr}*IF($A{r}>=31,1+{AUG}*'설정'!$E${tr},1)/$D{r})"
            lim = f"({BT}+'설정'!$F${tr}" + (f"+'설정'!$G${tr})" if boss else ")")
            sec = f"$E{r}*{C}*IF({R}<1,{R}^(-{LO}),{R}^(-{HI}))"
            ws.cell(r, col, f'=IF({sec}>{lim},"실패",ROUND({sec},0))')
            ws.cell(r, col).alignment = Alignment(horizontal='center')
    last = 2 + b.LAST; rng = f"H3:{L(7 + 16)}{last}"
    # 조건부 서식: 실패 / 85%↑ / 40%↓ (제한시간은 그 열 훈장 값)
    ws.conditional_formatting.add(rng, FormulaRule(formula=['H3="실패"'], fill=PatternFill('solid', fgColor='F4B6B6')))
    lim_rel = f"({BT}+INDEX('설정'!$F${T_ROW0 + 1}:$F${T_ROW0 + 16},COLUMN(H3)-7)+IF($C3=\"보스\",INDEX('설정'!$G${T_ROW0 + 1}:$G${T_ROW0 + 16},COLUMN(H3)-7),0))"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER(H3),H3>=0.85*{lim_rel})'], fill=PatternFill('solid', fgColor='FCE4B6')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER(H3),H3<0.4*{lim_rel})'], fill=PatternFill('solid', fgColor='D7EFD2')))
    ws.freeze_panes = 'H3'
    for c, w in enumerate([6, 8, 6, 15, 9, 13, 11] + [8] * 16, 1): ws.column_dimensions[L(c)].width = w


def simulate(p, tiers):
    for k in ('BASE_TIME', 'TIME_C', 'TIME_LO', 'TIME_HI', 'PASSIVE_SLOW', 'FLOOR_OVERHEAD', 'RUN_OVERHEAD', 'RAMP_SEC', 'LATE_SHARE', 'AUG_PER_SLOT', 'BIRTH_PER_POP'):
        setattr(b, k, float(p[k]))
    for t, r in b.TIER.items(): r['research_cost'] = r['research_cost'] * tiers[t]['cost']
    for s in b.SKILL:
        m = tiers[s['tier']]['cost']; s['cost_research'] *= m; s['cost_cheese'] *= m
    power = {str(t): dict(start=v['start'], end=v['end'], pop=v['pop']) for t, v in tiers.items()}
    out = {}
    for passive in (False, True):
        pl = b.Player(power, passive); rows = []; stuck = collections.Counter(); orig = pl.run; clear_at = {}
        def run():
            t0 = pl.t; s0 = pl.start_floor(); tier = pl.tier; tf = pl.tree_frac(); f = orig()
            rows.append([pl.runs, round(t0 / 3600, 2), tier, round(tf, 2), s0, f - 1, f if f <= b.LAST else '', round((pl.t - t0) / 60, 1), pl.res, pl.chz, pl.shard_level_sum(), pl.dex()])
            if f in HURDLES: stuck[f] += 1
            for m in MARKS:
                if m not in clear_at and pl.maxf - 1 >= m: clear_at[m] = pl.t
            return f
        pl.run = run; pl.play(hours_cap=float(p['HOURS_CAP']))
        out[passive] = dict(total=pl.t / 3600, runs=pl.runs, tier=pl.tier, maxf=pl.maxf - 1, rows=rows, stuck=stuck, clear=clear_at)
    return out


def write_results(wb, res, p):
    for n in ('요약', '판_클릭', '판_미클릭'):
        if n in wb.sheetnames: del wb[n]
    ws = wb.create_sheet('요약', 1)
    ws['A1'] = '시뮬레이션 결과 (이 시트는 스크립트가 덮어씀)'; ws['A1'].font = Font(bold=True, size=13)
    hd = ['', '클릭 적극', '거의 안 함']
    for c, h in enumerate(hd, 1):
        x = ws.cell(3, c, h); x.fill = HEAD; x.font = BOLD
    A, P = res[False], res[True]
    lines = [('총 시간 (h)', A['total'], P['total']), ('목표 (h)', p['GOAL_ACTIVE'], p['GOAL_PASSIVE']), ('판 수', A['runs'], P['runs']),
             ('판 평균 (분)', A['total'] * 60 / max(1, A['runs']), P['total'] * 60 / max(1, P['runs'])), ('최고 층', A['maxf'], P['maxf']), ('마지막 훈장', A['tier'], P['tier'])]
    lines += [(f'{m}층 클리어 (h)', A['clear'].get(m, 0) / 3600 or '—', P['clear'].get(m, 0) / 3600 or '—') for m in MARKS]
    lines += [(f'{h}층에서 막힌 판', A['stuck'].get(h, 0), P['stuck'].get(h, 0)) for h in HURDLES]
    for i, row in enumerate(lines):
        for c, v in enumerate(row, 1):
            x = ws.cell(4 + i, c, round(v, 2) if isinstance(v, float) else v)
    ws.cell(4 + len(lines) + 1, 1, '시간표 목표(누적, 사용자 2026-10-19): 튜토리얼 ~20분 · 16층 2h · 30층 4h · 31층 증강 4h · 50층 6.5h · 70층 9h · 100층 12.5h · 허들 층 2~3판')
    ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 14
    for passive, name in ((False, '판_클릭'), (True, '판_미클릭')):
        w = wb.create_sheet(name)
        for c, h in enumerate(['판', '시작 시각 (h)', '훈장', '트리 비율', '시작 층', '최고 클리어', '실패 층', '걸린 분', '남은 연구자료', '남은 치즈', '조각 합', '도감'], 1):
            x = w.cell(1, c, h); x.fill = HEAD; x.font = BOLD
        for i, row in enumerate(res[passive]['rows'], 2):
            for c, v in enumerate(row, 1):
                x = w.cell(i, c, v)
                if c == 4: x.number_format = '0%'
                if c in (9, 10): x.number_format = '0.00E+00'
        w.freeze_panes = 'A2'


def main():
    if os.path.exists(OUT):
        wb = openpyxl.load_workbook(OUT)
        if '설정' in wb.sheetnames and wb['설정']['Z1'].value != LAYOUT: del wb['설정']        # 칸 배치가 바뀐 옛 설정
        if '설정' not in wb.sheetnames: make_settings(wb, tier_static())
    else:
        wb = openpyxl.Workbook(); del wb[wb.sheetnames[0]]; make_settings(wb, tier_static())
    st = tier_static()
    ws = wb['설정']
    for t in range(1, 17):                       # 테이블에서 오는 칸은 매번 새로 (노란 칸은 그대로)
        r = T_ROW0 + t
        ws.cell(r, 6, st[t]['ta']); ws.cell(r, 7, st[t]['bt']); ws.cell(r, 8, st[t]['cost']).number_format = '#,##0'; ws.cell(r, 9, st[t]['floor'])
    p, tiers = read_settings(ws)
    floor_sheet(wb, st)
    res = simulate(p, tiers)
    write_results(wb, res, p)
    wb.save(OUT)
    for passive in (False, True):
        r = res[passive]
        print('거의 안 함' if passive else '클릭 적극', f"총 {r['total']:.1f}h · 판 {r['runs']} · 최고 {r['maxf']}층 ·",
              ' · '.join(f"{m}층 {r['clear'][m] / 3600:.1f}h" for m in MARKS if m in r['clear']))
    print('→', OUT)


if __name__ == '__main__':
    main()
