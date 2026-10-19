"""보스 체력 계수(보스 테이블 Boss 시트 hp_pow_sec) 를 지금 스테이지 테이블 보스전 목표 초(boss_sec)에 맞춤 (2026-10-19 전체 밸런스).
보스 체력 = 적정 전투력 × hp_pow_sec (Boss.HpFor). 아래 BASE = 예전(제한시간 180초 · 층 유형 반복 시절) 값과 그때 보스전 목표 초.
새 값 = BASE × (지금 boss_sec ÷ 예전 보스전 목표 초) — 몇 번 돌려도 같은 결과.
측정으로 보스전 시간이 목표와 다르면 BASE 를 고칠 것 (update_spu.py 가 보스전 비율을 알려 줌).
순서: gen_stage_table → … → gen_expansion_data → (이 스크립트) → xlsx2json
사용: python Tools/gen_boss_hp.py"""
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOSS = os.path.join(ROOT, 'Data_Table', '보스 테이블.xlsx')
STAGE = os.path.join(ROOT, 'Data_Table', '스테이지 테이블.xlsx')

# 층: (예전 hp_pow_sec, 예전 보스전 목표 초 = 0.4 × 예전 층 유형 목표: 허들 162 · 보통 108 · 어려움 135)
BASE = {5: (30.1, 64.8), 10: (33.6, 43.2), 15: (42.4, 43.2), 20: (50.5, 64.8), 25: (45.0, 64.8), 30: (45.0, 64.8), 35: (45.0, 64.8),
        40: (45.0, 64.8), 45: (48.0, 64.8), 50: (50.0, 64.8), 55: (48.0, 64.8), 60: (52.0, 64.8), 65: (48.0, 64.8), 70: (50.0, 64.8),
        75: (52.0, 64.8), 80: (50.0, 64.8), 85: (52.0, 64.8), 90: (55.0, 64.8), 94: (50.0, 54.0), 97: (52.0, 54.0), 100: (65.0, 64.8)}


def main():
    st = openpyxl.load_workbook(STAGE, data_only=True)['Stage']
    K = [c.value for c in st[2]]
    boss_sec = {r[0]: r[K.index('boss_sec')] for r in st.iter_rows(min_row=4, values_only=True) if isinstance(r[0], int)}
    wb = openpyxl.load_workbook(BOSS); ws = wb['Boss']
    keys = [c.value for c in ws[2]]; cf, ch = keys.index('floor') + 1, keys.index('hp_pow_sec') + 1
    for r in range(4, ws.max_row + 1):
        f = ws.cell(r, cf).value
        if not isinstance(f, int) or f not in BASE or not boss_sec.get(f): continue
        hp0, sec0 = BASE[f]
        v = round(hp0 * boss_sec[f] / sec0, 1)
        print(f'{f}층 hp_pow_sec {ws.cell(r, ch).value} → {v} (보스전 목표 {boss_sec[f]}초)')
        ws.cell(r, ch).value = v
    wb.save(BOSS)


if __name__ == '__main__':
    main()
