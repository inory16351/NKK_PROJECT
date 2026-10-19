"""조각 성장 비용·지급량 (2026-10-19 전체 밸런스, 사용자: '레벨업 조각 비용을 올림').
출생마다 그 종 조각 +1 이라 조각 합이 훈장 조건(Shard_Level_Sum)의 ~4배로 쌓였음 (balance_sim, 층당 출생 ≈ 최대 마릿수).
→ 레벨업 조각(쥐 등급 테이블 Grade.shard_k) × SHARD_SCALE, 뽑기(Gacha_Grade.shard_count)·암시장(Shop_Goods.pack_size) 조각도 같은 배율.
   ×10 = 승급할 때 조각 합 ≈ 조건의 1.1~1.4배 (×12 부터 조각 때문에 기다리는 판이 생김). 아래 BASE 는 예전 값이라 몇 번 돌려도 같은 결과.
사용: python Tools/gen_shards.py && python Tools/xlsx2json.py"""
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DT = os.path.join(ROOT, 'Data_Table')
SHARD_SCALE = 10
BASE_K = {'Common': 2.0, 'Rare': 1.5, 'Epic': 1.0, 'Unique': 0.8, 'Legendary': 0.6, 'Mythic': 0.5}
BASE_GACHA = {'Common': 1, 'Rare': 1, 'Epic': 1, 'Unique': 1, 'Legendary': 1, 'Mythic': 1}
BASE_PACK = {90001: 3, 90002: 6, 90003: 10, 90011: 2, 90012: 4, 90013: 6}


def edit(book, sheet, key, col, table):
    p = os.path.join(DT, book + '.xlsx'); wb = openpyxl.load_workbook(p); ws = wb[sheet]
    K = [c.value for c in ws[2]]; ki, ci = K.index(key) + 1, K.index(col) + 1
    for r in range(4, ws.max_row + 1):
        k = ws.cell(r, ki).value
        if k in table:
            v = table[k] * SHARD_SCALE
            ws.cell(r, ci, int(v) if float(v).is_integer() else v)
    wb.save(p); print(book, sheet, col, {k: v * SHARD_SCALE for k, v in table.items()})


def main():
    edit('쥐 등급 테이블', 'Grade', 'grade', 'shard_k', BASE_K)
    edit('뽑기 테이블', 'Gacha_Grade', 'grade', 'shard_count', BASE_GACHA)
    edit('상점 테이블', 'Shop_Goods', 'goods_id', 'pack_size', BASE_PACK)


if __name__ == '__main__':
    main()
