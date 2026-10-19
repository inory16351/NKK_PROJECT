"""100층 확장 데이터 → 엑셀 (2026-10-19). 내용은 Tools/expansion_content.py.
물건 테이블 (Zone · Item · Furniture_Layout) · 사람 테이블 (Human · Human_Line) · 보스 테이블 (Boss · Boss_Line · Atk_Type) · 티어 테이블 (Tier · Cond_Type)
같은 id(Atk_Type 은 atk_type) 줄은 덮어쓰고, 없으면 끝에 붙임. 새 칸(Zone themes · Human to_floor · Atk_Type pattern… · Tier cond4)이 없으면 만듦.
고치기 전 원본은 Data_Table/_backup_20261019/ 에 한 번 복사.
순서: gen_stage_table → gen_economy → gen_expansion_data → gen_skill_tree → gen_string_table → (이 스크립트 --en: 스트링 테이블 영어 채움) → xlsx2json
사용: python Tools/gen_expansion_data.py [--en]"""
import copy, io, os, shutil, sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import expansion_content as C

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DT = os.path.join(ROOT, 'Data_Table')
BACKUP = os.path.join(DT, '_backup_20261019')


def book(name):
    p = os.path.join(DT, name + '.xlsx')
    os.makedirs(BACKUP, exist_ok=True)
    b = os.path.join(BACKUP, name + '.xlsx')
    if not os.path.exists(b): shutil.copy2(p, b)
    return openpyxl.load_workbook(p), p


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.keys = [c.value for c in ws[2]]
        while self.keys and self.keys[-1] is None: self.keys.pop()

    def col(self, key, ko=None, typ='string'):
        if key in self.keys: return self.keys.index(key) + 1
        c = len(self.keys) + 1
        for r, v in enumerate((ko or key, key, typ), 1):
            cell = self.ws.cell(r, c, v); cell._style = copy.copy(self.ws.cell(r, c - 1)._style)
        self.keys.append(key)
        return c

    def find(self, key, value):
        ci = self.keys.index(key) + 1
        for r in range(4, self.ws.max_row + 1):
            if self.ws.cell(r, ci).value == value: return r
        return None

    def last_row(self):
        r = self.ws.max_row
        while r > 3 and self.ws.cell(r, 1).value is None: r -= 1
        return r

    def upsert(self, key, value, data):
        r = self.find(key, value)
        if r is None:
            r = self.last_row() + 1
            for c in range(1, len(self.keys) + 1): self.ws.cell(r, c)._style = copy.copy(self.ws.cell(r - 1, c)._style)
        data = dict(data); data[key] = value
        for k, v in data.items():
            if k not in self.keys: continue
            self.ws.cell(r, self.keys.index(k) + 1, v)
        return r


EN = {}       # 한국어 → 영어 (스트링 테이블 채우기)
def ko(pair):
    k, e = pair; EN[k] = e; return k


def items_and_zones():
    wb, p = book('물건 테이블')
    z = Sheet(wb['Zone']); z.col('themes', '가구 테마 (쉼표)', 'string')
    for zid in (70001, 70002, 70003, 70004):
        r = z.find('zone_id', zid)
        if r: z.ws.cell(r, z.keys.index('themes') + 1, C.OLD_THEMES)
    for zid, name, f0, f1, themes in C.ZONES:
        z.upsert('zone_id', zid, dict(zone_name=ko(name), from_floor=f0, floor_type='Lab', tint=None, themes=themes))
    it = Sheet(wb['Item'])
    # 예전 연구소 일반 물건: 30층까지 (31층~ 은 구간 물건만)
    for r in range(4, it.ws.max_row + 1):
        iid = it.ws.cell(r, 1).value
        if isinstance(iid, int) and 60001 <= iid <= 60024 and not it.ws.cell(r, it.keys.index('to_floor') + 1).value:
            it.ws.cell(r, it.keys.index('to_floor') + 1, 30)
    fl = Sheet(wb['Furniture_Layout'])
    LAYOUT_A = [(0, 0.18, 0.10), (1, 0.52, 0.12), (3, 0.86, 0.12), (2, 0.30, 0.62), (4, 0.72, 0.64)]
    LAYOUT_B = [(1, 0.20, 0.12), (4, 0.60, 0.10), (0, 0.88, 0.14), (3, 0.40, 0.62), (2, 0.76, 0.70)]
    lid = 80100
    for zi, (zid, _, f0, f1, themes) in enumerate(C.ZONES, 1):
        rows = C.ZONE_ITEMS[zi]
        to = 0 if f1 >= 100 else f1
        normal = [62000 + zi * 100 + i for i in range(1, 11)]
        for i, (code, k, e) in enumerate(rows):
            n = i + 1
            if n <= 10:
                iid = 62000 + zi * 100 + n
                hp = C.ITEM_HP[i] * (1 + 0.15 * zi)
                d = dict(item_category='Item', radius=14 + (i % 4) * 1.5 + (4 if hp > 2 else 0), hp_mul=round(hp, 2), value_mul=round(hp * 1.05, 2), spawn_weight=1.2,
                         heavy=0, is_big=0, is_sturdy=1 if hp > 1.6 else 0, is_paper=1 if code in ('cat_bread', 'zero_days_sign', 'doge_poster', 'flex_poster', 'pointing_poster', 'market_price') else 0,
                         spill_color=C.SPILL_COLOR[zi] if i % 3 == 0 else None, drop_01=0, drop_02=0, drop_03=0, drop_04=0)
            elif n <= 13:
                iid = 62000 + zi * 100 + n
                d = dict(item_category='Item', radius=24, hp_mul=round((7 + (n - 11) * 3) * (1 + 0.1 * zi), 1), value_mul=round((12 + (n - 11) * 6) * (1 + 0.1 * zi), 1), spawn_weight=0.35,
                         heavy=0, is_big=1 if n == 13 else 0, is_sturdy=1, is_paper=0, spill_color=None,
                         drop_01=60033, drop_02=60034, drop_03=60035, drop_04=60036 if n != 12 else 60035)
            else:
                iid = 63000 + zi * 100 + (n - 13)
                fi = n - 14
                d = dict(item_category='Furniture', radius=C.FURN_RAD[fi], hp_mul=round(16 + fi * 2 + zi, 1), value_mul=round(11 + fi * 2 + zi, 1), spawn_weight=0,
                         heavy=4 + (fi % 3), is_big=1, is_sturdy=1, is_paper=0, spill_color=None,
                         drop_01=normal[(fi * 2) % 10], drop_02=normal[(fi * 2 + 3) % 10], drop_03=normal[(fi * 2 + 7) % 10], drop_04=0)
            d.update(item_name=ko((k, e)), code_id=code, asset=f'Expansion/Items/{code}',
                     from_floor=0 if n > 13 else f0, to_floor=0 if n > 13 else to, unlock_skill=0)
            it.upsert('item_id', iid, d)
        for t, lay in ((themes.split(',')[0], LAYOUT_A), (themes.split(',')[1], LAYOUT_B)):
            for fi, u, v in lay:
                lid += 1
                fl.upsert('layout_id', lid, dict(room_theme=t, item_id=63000 + zi * 100 + fi + 1, pos_u=u, pos_v=v))
    wb.save(p)
    print('물건 · 구간 · 가구 배치')


def humans():
    wb, p = book('사람 테이블')
    h = Sheet(wb['Human']); h.col('to_floor', '사라지는 층 (0 = 계속)', 'int')
    for hid, end in C.OLD_HUMAN_END.items():
        r = h.find('human_id', hid)
        if r: h.ws.cell(r, h.keys.index('to_floor') + 1, end)
    hl = Sheet(wb['Human_Line'])
    lid = 91100
    for zi, (zid, _, f0, f1, _) in enumerate(C.ZONES, 1):
        for i, (code, k, e, hp, val, spd, lines) in enumerate(C.ZONE_HUMANS[zi]):
            hid = 90010 + (zi - 1) * 4 + i + 1
            h.upsert('human_id', hid, dict(human_name=ko((k, e)), code_id=code, hp_mul=hp, value_mul=val, speed_mul=spd, from_floor=f0,
                                           to_floor=0 if f1 >= 100 else f1, asset_folder=f'Humans/Parts/{code}'))
            for ln in lines:
                lid += 1
                hl.upsert('line_id', lid, dict(human_id=hid, situation='Panic', text=ko(ln)))
    wb.save(p)
    print('사람', len(C.ZONE_HUMANS) * 4)


def bosses():
    wb, p = book('보스 테이블')
    at = Sheet(wb['Atk_Type'])
    for key, kn, typ in [('pattern', '밈 패턴 (Boss.Meme.cs)', 'enum'), ('sprite', '그림 (BossMeme)', 'string'), ('size', '그림 크기', 'float'),
                         ('color', '색', 'string'), ('popup', '팝업 글', 'string'), ('p1', '패턴 값 1', 'float'), ('p2', '패턴 값 2', 'float')]:
        at.col(key, kn, typ)
    for (t, name, pat, rad, stun, cnt, wind, dur, spr, size, color, pop, p1, p2, desc) in C.MEMES:
        at.upsert('atk_type', t, dict(desc=desc, radius=rad, stun=stun, count=cnt, windup=wind, dur=dur, atk_name=ko(name), line=None,
                                      pattern=pat, sprite=spr, size=size, color=color, popup=ko(pop), p1=p1, p2=p2))
    b = Sheet(wb['Boss'])
    for bid in C.RETIRED_BOSSES:
        r = b.find('boss_id', bid)
        if r:
            b.ws.cell(r, b.keys.index('floor') + 1, 0)
            dc = b.keys.index('-') + 1 if '-' in b.keys else None
            if dc:
                v = b.ws.cell(r, dc).value or ''
                if '미사용' not in v: b.ws.cell(r, dc, '(미사용 2026-10-19: 16~30층 고양이 연구소로 교체, 테스트 패널로만) ' + v)
    tmpl = b.find('boss_id', 95004)
    base = {k: b.ws.cell(tmpl, i + 1).value for i, k in enumerate(b.keys)}
    for (bid, name, code, rig, floor, color, scale, a1, a2, sk, s1, s2, hp) in C.BOSSES:
        d = dict(base)
        d.update(boss_name=ko(name), code_id=code, rig=rig, floor=floor, atk_type=a1, atk2_type=a2, atk2_chance=0.4, color=color, scale=scale, radius_mul=scale,
                 hp_pow_sec=hp, skill_type=sk, skill_cd=12, skill_first=6, special1_type=s1, special2_type=s2, special_chance=0.25)
        d['-'] = f'{floor}층 보스 (2026-10-19 100층 확장). 기술 = Atk_Type {sk} · {s1} · {s2}'
        d.pop('boss_id', None); d.pop('boss_name_sk', None)
        b.upsert('boss_id', bid, d)
    bl = Sheet(wb['Boss_Line'])
    lid = 96200
    for bid, sit in C.BOSS_LINES.items():
        for s, lines in sit.items():
            for ln in lines:
                lid += 1
                bl.upsert('line_id', lid, dict(boss_id=bid, situation=s, text=ko(ln)))
    wb.save(p)
    print('보스', len(C.BOSSES), '· 밈 기술', len(C.MEMES))


LATE_COST_SCALE = 1.95       # 9~16훈장 승급 비용 배율 (2026-10-19, balance_sim 으로 맞춤 · 트리 비용은 gen_skill_tree 가 다음 승급 비용에서 따라감)


def tiers():
    st = openpyxl.load_workbook(os.path.join(DT, '스테이지 테이블.xlsx'), data_only=True)['Stage']
    keys = [c.value for c in st[2]]
    research = {r[0]: r[keys.index('research')] for r in st.iter_rows(min_row=4, values_only=True) if isinstance(r[0], int)}
    wb, p = book('티어 테이블')
    t = Sheet(wb['Tier'])
    rc = t.keys.index('research_cost') + 1
    t.ws.cell(3, rc, 'float')
    t.col('cond4_type', '조건4 타입', 'enum'); t.col('cond4_value', '조건4 값', 'float')
    for (tier, name, icon, conds, bk, start, atk) in C.TIERS:
        floor = next((v for k, v in conds if k in ('Max_Floor', 'Boss_Clear')), 30)
        cost = round(research.get(min(floor, 100), 0) * 4 * LATE_COST_SCALE, -3)
        d = dict(tier_name=ko(name), icon=icon, research_cost=cost, birth_grade_k=bk, start_rat_count=start, start_floor_cap=0, badge_asset=f'Rogue/rk_{tier}', rat_atk_mul=atk)
        for i in range(4):
            k, v = conds[i] if i < len(conds) else ('None', 0)
            d[f'cond{i + 1}_type'] = k; d[f'cond{i + 1}_value'] = v
        t.upsert('tier', tier, d)
    ct = Sheet(wb['Cond_Type'])
    for k, desc, val in C.COND_DOC: ct.upsert('cond_type', k, dict(desc=desc, value=val))
    wb.save(p)
    print('훈장 9~16')


def fill_english():
    p = os.path.join(DT, '스트링 테이블.xlsx')
    wb = openpyxl.load_workbook(p); ws = wb['String']
    keys = [c.value for c in ws[2]]
    ki, ei, si = keys.index('ko') + 1, keys.index('en') + 1, keys.index('status') + 1
    n = 0
    for r in range(4, ws.max_row + 1):
        k = ws.cell(r, ki).value
        if k in EN and (not ws.cell(r, ei).value or ws.cell(r, si).value in ('번역 필요', '한글 바뀜')):
            ws.cell(r, ei, EN[k]); ws.cell(r, si, '번역 완료'); n += 1
    wb.save(p)
    print('영어 채움', n)


def main():
    items_and_zones(); humans(); bosses(); tiers()
    if '--en' in sys.argv: fill_english()


if __name__ == '__main__':
    main()
